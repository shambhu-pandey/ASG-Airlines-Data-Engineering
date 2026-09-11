"""Batch ingestion of the source workbook into the MySQL raw layer."""

import argparse
import hashlib
import json
import uuid
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text

from .db import get_engine


BATCH_SIZE = 200

DATASET_TABLES = {
	"flights": ("raw_flights", "raw_flight_id", [
		"flight_id", "airline", "source", "destination", "departure_time",
		"arrival_time", "duration",
	]),
	"bookings": ("raw_bookings", "raw_booking_id", [
		"booking_id", "passenger_id", "flight_id", "booking_date", "status",
		"passport_number", "seat_number", "emergency_contact_name",
		"emergency_contact_phone",
	]),
	"passengers": ("raw_passengers", "raw_passenger_id", [
		"passenger_id", "first_name", "last_name", "age", "gender", "email",
		"phone", "aadhaar_id", "date_of_birth",
	]),
	"payments": ("raw_payments", "raw_payment_row_id", [
		"payment_id", "booking_id", "amount", "payment_method",
	]),
}


def _json_value(value: Any) -> Any:
	"""Convert workbook values to JSON-safe values without business cleaning."""
	if isinstance(value, (datetime, date)):
		return value.isoformat(sep=" ")
	return value


def _record_hash(row: dict[str, Any]) -> str:
	serialized = json.dumps(row, sort_keys=True, ensure_ascii=False, default=str)
	return hashlib.sha256(serialized.encode("utf-8")).hexdigest()


def _file_hash(source_path: Path) -> str:
	digest = hashlib.sha256()
	with source_path.open("rb") as source_file:
		for chunk in iter(lambda: source_file.read(1024 * 1024), b""):
			digest.update(chunk)
	return digest.hexdigest()


def _source_file_name(source_path: Path) -> str:
	try:
		return source_path.resolve().relative_to(Path.cwd().resolve()).as_posix()
	except ValueError:
		return source_path.name


def _find_sheets(workbook: Any) -> dict[str, Any]:
	sheets = {sheet.title.strip().lower(): sheet for sheet in workbook.worksheets}
	missing = sorted(set(DATASET_TABLES) - set(sheets))
	if missing:
		raise ValueError(f"Workbook is missing expected sheets: {', '.join(missing)}")
	return {name: sheets[name] for name in DATASET_TABLES}


def _read_sheet_rows(sheet: Any, columns: list[str]) -> list[tuple[int, dict[str, Any]]]:
	header_cells = next(sheet.iter_rows(min_row=1, max_row=1))
	header_map = {
		cell.value.strip().lower(): index
		for index, cell in enumerate(header_cells)
		if isinstance(cell.value, str) and cell.value.strip()
	}
	missing = [column for column in columns if column not in header_map]
	if missing:
		raise ValueError(f"Sheet '{sheet.title}' is missing columns: {', '.join(missing)}")

	rows = []
	for excel_row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
		values = {
			column: _json_value(row[header_map[column]].value)
			for column in columns
		}
		rows.append((excel_row_number, values))
	return rows


def _insert_batch(
	connection: Any,
	table_name: str,
	rows: list[dict[str, Any]],
) -> None:
	if not rows:
		return
	columns = list(rows[0])
	column_sql = ", ".join(f"`{column}`" for column in columns)
	value_sql = ", ".join(f":{column}" for column in columns)
	connection.execute(
		text(f"INSERT INTO {table_name} ({column_sql}) VALUES ({value_sql})"),
		rows,
	)


def ingest_workbook(source_path: str | Path, batch_size: int = BATCH_SIZE) -> dict[str, Any]:
	"""Ingest workbook rows into raw tables and return a run summary."""
	path = Path(source_path)
	if not path.is_file():
		raise FileNotFoundError(f"Source workbook not found: {path}")
	if batch_size <= 0:
		raise ValueError("batch_size must be greater than zero")

	file_hash = _file_hash(path)
	source_file = _source_file_name(path)
	engine = get_engine()
	started_at = datetime.now(timezone.utc).replace(tzinfo=None)
	run_id = uuid.uuid4().hex
	workbook = load_workbook(path, read_only=True, data_only=False)

	try:
		with engine.begin() as connection:
			existing_run = connection.execute(
				text(
					"SELECT run_id FROM etl_runs "
					"WHERE file_hash = :file_hash AND status = 'succeeded' "
					"ORDER BY started_at DESC LIMIT 1"
				),
				{"file_hash": file_hash},
			).scalar_one_or_none()
			if existing_run:
				return {
					"run_id": existing_run,
					"status": "skipped",
					"reason": "source file hash was already ingested successfully",
				}

			connection.execute(
				text(
					"INSERT INTO etl_runs "
					"(run_id, source_file, file_hash, started_at, status) "
					"VALUES (:run_id, :source_file, :file_hash, :started_at, 'running')"
				),
				{
					"run_id": run_id,
					"source_file": source_file,
					"file_hash": file_hash,
					"started_at": started_at,
				},
			)

			sheets = _find_sheets(workbook)
			batch_count = 0
			next_batch_number = 1
			inserted_counts = {name: 0 for name in DATASET_TABLES}
			error_messages = []

			for dataset_name, sheet in sheets.items():
				table_name, _, columns = DATASET_TABLES[dataset_name]
				try:
					rows = _read_sheet_rows(sheet, columns)
				except Exception as error:
					raise RuntimeError(f"Unable to prepare sheet '{sheet.title}': {error}") from error

				for start in range(0, len(rows), batch_size):
					batch_rows = rows[start:start + batch_size]
					batch_number = next_batch_number
					next_batch_number += 1
					batch_id = uuid.uuid4().hex
					batch_started_at = datetime.now(timezone.utc).replace(tzinfo=None)
					prepared_rows = []
					row_errors = []

					for source_row_num, original_values in batch_rows:
						try:
							prepared_rows.append({
								"run_id": run_id,
								"batch_id": batch_id,
								"source_file": source_file,
								"source_row_num": source_row_num,
								"ingestion_timestamp": batch_started_at,
								"record_hash": _record_hash(original_values),
								**original_values,
								"raw_payload": json.dumps(original_values, ensure_ascii=False, default=str),
							})
						except Exception as error:
							row_errors.append(f"row {source_row_num}: {error}")

					batch_status = "succeeded" if not row_errors else "completed_with_errors"
					batch_error = "; ".join(row_errors)[:65000] or None
					connection.execute(
						text(
							"INSERT INTO etl_batches "
							"(batch_id, run_id, table_name, batch_number, source_row_start, "
							"source_row_end, row_count, status, started_at, ended_at, error_message) "
							"VALUES (:batch_id, :run_id, :table_name, :batch_number, :source_row_start, "
							":source_row_end, :row_count, :status, :started_at, :ended_at, :error_message)"
						),
						{
							"batch_id": batch_id,
							"run_id": run_id,
							"table_name": table_name,
							"batch_number": batch_number,
							"source_row_start": batch_rows[0][0],
							"source_row_end": batch_rows[-1][0],
							"row_count": len(prepared_rows),
							"status": batch_status,
							"started_at": batch_started_at,
							"ended_at": datetime.now(timezone.utc).replace(tzinfo=None),
							"error_message": batch_error,
						},
					)
					_insert_batch(connection, table_name, prepared_rows)
					batch_count += 1
					inserted_counts[dataset_name] += len(prepared_rows)
					if row_errors:
						error_messages.extend(f"{dataset_name} {message}" for message in row_errors)

			final_status = "completed_with_errors" if error_messages else "succeeded"
			connection.execute(
				text(
					"UPDATE etl_runs SET ended_at = :ended_at, status = :status, "
					"error_message = :error_message WHERE run_id = :run_id"
				),
				{
					"ended_at": datetime.now(timezone.utc).replace(tzinfo=None),
					"status": final_status,
					"error_message": "; ".join(error_messages)[:65000] or None,
					"run_id": run_id,
				},
			)
		return {
			"run_id": run_id,
			"status": final_status,
			"batch_count": batch_count,
			"inserted_counts": inserted_counts,
			"source_file": source_file,
			"file_hash": file_hash,
		}
	finally:
		workbook.close()


if __name__ == "__main__":
	parser = argparse.ArgumentParser(description="Ingest the ASG Airlines workbook into MySQL raw tables.")
	parser.add_argument("source", nargs="?", default="data/raw/UseCase - Airlines.xlsx")
	parser.add_argument("--batch-size", type=int, default=BATCH_SIZE)
	args = parser.parse_args()
	print(json.dumps(ingest_workbook(args.source, args.batch_size), indent=2))
